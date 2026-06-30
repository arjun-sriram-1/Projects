from pathlib import Path

from docx import Document

from api.core.config import settings


def _resolve_output_path(filename):
    output_path = Path(filename) if filename is not None else settings.report_dir / "credit_memo.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    return output_path


def export_docx(memo_text, filename=None):
    output_path = _resolve_output_path(filename)
    document = Document()
    document.add_heading("Credit Committee Memo", level=1)
    document.add_paragraph(memo_text)
    document.save(str(output_path))
    return str(output_path)
